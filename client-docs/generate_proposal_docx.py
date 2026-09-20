from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

doc = Document()

for section in doc.sections:
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

NAVY = RGBColor(0x0F, 0x3A, 0x47)
TEAL = RGBColor(0x1F, 0x6B, 0x5C)
BODY = RGBColor(0x2C, 0x2C, 0x2C)
MUTED = RGBColor(0x5A, 0x5A, 0x5A)
SOFT = RGBColor(0xD5, 0xE8, 0xE2)


def set_run_font(run, name='Calibri', size=11, bold=False, color=BODY, italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def add_para(text, size=11, bold=False, color=BODY, space_before=0, space_after=8, align='left', italic=False):
    p = doc.add_paragraph()
    p.alignment = {
        'left': WD_ALIGN_PARAGRAPH.LEFT,
        'center': WD_ALIGN_PARAGRAPH.CENTER,
        'justify': WD_ALIGN_PARAGRAPH.JUSTIFY,
    }[align]
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color, italic=italic)
    return p


def shade_cell(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)


def set_cell_borders(cell, color='C5D5CF'):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = parse_xml(
        f'''<w:tcBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:left w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
        </w:tcBorders>'''
    )
    tcPr.append(borders)


def write_cell(cell, text, bold=False, size=10, color=BODY, fill=None):
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    if fill:
        shade_cell(cell, fill)
    set_cell_borders(cell)


def add_section_title(number, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(f'{number}  {title}')
    set_run_font(run, size=13, bold=True, color=NAVY)
    pPr = p._p.get_or_add_pPr()
    pBdr = parse_xml(
        f'''<w:pBdr {nsdecls("w")}>
            <w:bottom w:val="single" w:sz="12" w:space="4" w:color="1F6B5C"/>
        </w:pBdr>'''
    )
    pPr.append(pBdr)


def add_subhead(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_run_font(run, size=11, bold=True, color=TEAL)


def add_bullet(text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Cm(0.75)
    if p.runs:
        p.runs[0].text = text
        set_run_font(p.runs[0], size=10.5, color=BODY)
    else:
        run = p.add_run(text)
        set_run_font(run, size=10.5, color=BODY)


def add_numbered(text, n):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Cm(0.4)
    run = p.add_run(f'{n}.  {text}')
    set_run_font(run, size=10.5, color=BODY)


def build_cap(cell, title, bullets, fill='F4F7F6'):
    cell.text = ''
    shade_cell(cell, fill)
    set_cell_borders(cell, '1F6B5C')
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(title)
    set_run_font(r, size=11, bold=True, color=TEAL)
    for b in bullets:
        bp = cell.add_paragraph()
        bp.paragraph_format.space_before = Pt(1)
        bp.paragraph_format.space_after = Pt(3)
        bp.paragraph_format.left_indent = Cm(0.15)
        br = bp.add_run('•  ' + b)
        set_run_font(br, size=9.5, color=BODY)
    end = cell.add_paragraph()
    end.paragraph_format.space_after = Pt(6)


# Header banner
banner = doc.add_table(rows=1, cols=1)
bc = banner.cell(0, 0)
bc.text = ''
bp = bc.paragraphs[0]
r1 = bp.add_run('SMARTGREENECOS  ·  PLATFORM 2')
set_run_font(r1, size=9, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
bp2 = bc.add_paragraph()
r2 = bp2.add_run('Delivery Proposal')
set_run_font(r2, size=22, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
bp3 = bc.add_paragraph()
r3 = bp3.add_run('Capacity building for dairy SMEs  ·  Training course + 9 modules')
set_run_font(r3, size=9.5, color=SOFT)
shade_cell(bc, '0F3A47')
set_cell_borders(bc, '0F3A47')
for p in bc.paragraphs:
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
bc.paragraphs[0].paragraph_format.space_before = Pt(10)
bc.paragraphs[-1].paragraph_format.space_after = Pt(10)

doc.add_paragraph().paragraph_format.space_after = Pt(4)

meta = doc.add_table(rows=4, cols=2)
meta_rows = [
    ('To', '[Mother Company Name]'),
    ('From', 'Investment World for Development and Technology'),
    ('Date', 'August 2026'),
    ('Subject', 'Platform 2 — capacity building for dairy SMEs'),
]
for i, (k, v) in enumerate(meta_rows):
    write_cell(meta.cell(i, 0), k, bold=True, size=10, color=RGBColor(0xFF, 0xFF, 0xFF), fill='1F6B5C')
    write_cell(meta.cell(i, 1), v, bold=False, size=10, color=BODY, fill='F4F7F6')
for row in meta.rows:
    row.cells[0].width = Cm(3.2)
    row.cells[1].width = Cm(13.5)

add_section_title('01', 'Purpose')
add_para(
    'This is what we will build for Platform 2: a capacity-building site for people working in dairy companies.',
    size=10.5, align='justify', space_after=6,
)
add_para(
    'It includes one training course and nine learning modules. Modules 2 to 10 each have a matching agent. Item 1 is a training course only (using AI platforms) — materials live inside that course; we do not build a separate agent for it.',
    size=10.5, align='justify', space_after=4,
)

add_section_title('02', 'What we are building')
add_para('There are two ways to use the platform.', size=10.5, space_after=6)

add_subhead('A. Training course (item 1) and learning modules (2–10)')
add_para(
    'Someone new to AI starts with the training course (materials, examples, exercises, test). For operational topics they go through a module: objectives, practice company, questions, agent, test, then upload. Section 3 spells this out.',
    size=10.5, align='justify', space_after=6,
)

add_subhead('B. Agents')
add_para(
    'If they already have their numbers in a spreadsheet, they can skip the module path. They open the agent, upload the file, and run the analysis. It is the same agent as in modules 2–10.',
    size=10.5, align='justify', space_after=6,
)

add_section_title('03', 'What the platform will cover')
add_subhead('3.1  Training course — Item 1 (materials inside the course)')
add_para(
    'Item 1 is a training course, not a module with an agent. All teaching material for “how to use AI platforms and design simple agents” sits inside this course — short lessons, examples, guided exercises, and a knowledge check. No agent. No company file upload.',
    size=10.5, align='justify', space_after=4,
)
add_para(
    'Course materials: learning objectives; short lesson content (what an agent is, what data to prepare, how to read outputs); worked dairy examples; guided exercises with feedback; knowledge check with a score. Typical time: about 45–90 minutes.',
    size=10.5, align='justify', space_after=6,
)

add_subhead('3.2  What is inside each learning module (2–10)')
add_para(
    'Every module from 2 to 10 follows the same steps. The topic, the practice table, the questions, and the agent change.',
    size=10.5, align='justify', space_after=6,
)

add_subhead('1. Learning objectives')
add_para(
    'A few lines on what they should be able to do afterwards — for example, read sales by product, make sense of a forecast, and use it so they waste less or run out less often.',
    size=10.5, align='justify', space_after=4,
)

add_subhead('2. Practice company data')
add_para(
    'They get a full example from a made-up dairy plant, as a table they can open. For demand forecasting that is eight weeks of weekly sales by product: name, category, units, and whether it went to shops or food service. They also see an external signals calendar — holidays, school terms, promotions, Ramadan — so they can link spikes in sales to real-world causes. Other modules use the same idea with different tables. We show a few headline numbers first so they know what to look for before they open the whole table.',
    size=10.5, align='justify', space_after=4,
)

add_subhead('3. Guided exercises')
add_para(
    'Questions about that table, one at a time. For forecasting, at least one question links a demand spike to an external signal (e.g. school term start in Week 5). They have to answer before they move on.',
    size=10.5, align='justify', space_after=4,
)

add_subhead('4. Run the agent')
add_para(
    'They run the agent on the practice data. For forecasting they see next week’s volumes by product, whether demand is up or down, which external signals the agent used (school term, promotion, holiday, Ramadan), what we suggest, and what could go wrong.',
    size=10.5, align='justify', space_after=4,
)

add_subhead('5. Short test')
add_para(
    'A few more questions, one at a time, to check they understood the idea. They get feedback after each one and a score at the end.',
    size=10.5, align='justify', space_after=4,
)

add_subhead('6. Use it on the company’s own files')
add_para(
    'They attach the kind of export they would pull from their own system (CSV). For forecasting, weekly sales by product is required; an external signals file (holidays, promotions, school terms, Ramadan) is optional but recommended. They can upload their file, download a blank template, or try a demo export if they have nothing ready. Then they run the same agent on that file and see results for that company — not the practice table from step 2.',
    size=10.5, align='justify', space_after=4,
)
add_para(
    'A full module takes about 45–90 minutes. The online demo is a shorter version of the same path.',
    size=10, color=MUTED, space_after=4, italic=True,
)

add_subhead('External signals in Module 2 (demand forecasting)')
add_para(
    'External signals are calendar events that explain demand spikes — not only sales numbers. For a Jordan dairy SME we tag: school term start/end (yogurt, breakfast milk), public holidays (shorter retail week), Ramadan and pre-Ramadan stock-up (labneh, cheese, UHT), retail promotions (promoted SKU lift), weather (chilled products), and food-service contracts.',
    size=10.5, align='justify', space_after=4,
)
add_para(
    'In Module 2 step 2 learners see two tables: weekly sales (8 weeks, 4 SKUs) and an external signals calendar — e.g. Week 5 school term start, Week 6 yogurt promotion, Week 9 local holiday. The exercise asks which signal explains the Week 5–6 yogurt spike. Agent output and company upload both include an “External signals used” block. Upload columns: week, event_name, event_type (school_term, holiday, promotion, ramadan, weather, other), optional expected_impact and notes. Sales CSV required; signals CSV optional but recommended.',
    size=10.5, align='justify', space_after=6,
)

add_subhead('3.3  Delivery list — training course + modules')
add_para(
    'Item 1 is the training course (materials only). Modules 2 to 10 are each one module and one agent, using the same files and the same kind of analysis.',
    size=10.5, align='justify', space_after=6,
)

mods = [
    '1  Using AI platforms and designing simple agents  ·  TRAINING COURSE (materials only)',
    '2  Demand forecasting with an AI forecasting agent  ·  Module + agent',
    '3  Sales and order analysis with an AI sales agent  ·  Module + agent',
    '4  Production planning and sequencing with an AI production-planning agent  ·  Module + agent',
    '5  Inventory monitoring with an AI inventory agent  ·  Module + agent',
    '6  Shelf-life management with an AI shelf-life agent  ·  Module + agent',
    '7  Milk procurement with an AI supplier agent  ·  Module + agent',
    '8  Predictive maintenance with an AI maintenance agent  ·  Module + agent',
    '9  Profitability and costing with an AI margin agent  ·  Module + agent',
    '10  Customer complaint analysis with an AI complaint agent  ·  Module + agent',
]
mt = doc.add_table(rows=1 + len(mods), cols=1)
write_cell(mt.cell(0, 0), 'Delivery item', bold=True, size=10, color=RGBColor(0xFF, 0xFF, 0xFF), fill='0F3A47')
for i, name in enumerate(mods, 1):
    fill = 'F4F7F6' if i % 2 else 'FFFFFF'
    write_cell(mt.cell(i, 0), name, bold=False, size=9.5, color=BODY, fill=fill)

add_para(
    'Note on Module 4: production sequencing (run order, changeovers) is part of this module — not a separate item.',
    size=10, color=MUTED, space_before=6, space_after=4, italic=True,
)
add_para(
    'Better forecasts and stock decisions also cut waste. That is how green practice shows up here — inside the day-to-day work, not as a separate tool.',
    size=10.5, align='justify', space_before=8, space_after=4,
)

add_subhead('3.4  What you can already try')
add_para(
    'Module 2 — Demand forecasting is live: objectives, eight-week sales table, external signals calendar, exercises, agent, test, and file upload. The forecasting agent is also available on its own for people who already have a file.',
    size=10.5, align='justify', space_after=4,
)
add_para(
    'Demo: https://muneernas.github.io/dairy-platform2-prototype/',
    size=10.5, space_after=4,
)
add_para(
    'The other modules will follow this same shape. Item 1 stays a training course without its own agent. All ten items are in this delivery.',
    size=10.5, align='justify', space_after=4,
)

add_section_title('04', 'Summary')
summary = [
    ('Purpose', 'Help dairy SMEs get better at using AI in everyday work'),
    ('How it works', 'One training course (item 1) + learning modules (2–10) with agents'),
    ('Inside training course', 'Materials, examples, exercises, knowledge check'),
    ('Inside a module', 'Objectives, practice table, external signals (where relevant), questions, agent, test, upload'),
    ('Scope', 'Training course (1) + modules 2–10; agents for modules 2–10'),
    ('Already working', 'Demand forecasting (Module 2)'),
]
st = doc.add_table(rows=1 + len(summary), cols=2)
write_cell(st.cell(0, 0), 'Topic', bold=True, size=10, color=RGBColor(0xFF, 0xFF, 0xFF), fill='0F3A47')
write_cell(st.cell(0, 1), 'What we will deliver', bold=True, size=10, color=RGBColor(0xFF, 0xFF, 0xFF), fill='0F3A47')
for i, (a, b) in enumerate(summary, 1):
    fill = 'F4F7F6' if i % 2 else 'FFFFFF'
    write_cell(st.cell(i, 0), a, bold=True, size=9.5, color=NAVY, fill=fill)
    write_cell(st.cell(i, 1), b, bold=False, size=9.5, color=BODY, fill=fill)
for row in st.rows:
    row.cells[0].width = Cm(4.0)
    row.cells[1].width = Cm(12.7)

fp = doc.add_paragraph()
fp.paragraph_format.space_before = Pt(18)
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = fp.add_run('Platform 2  ·  Capacity building  ·  Dairy sector')
set_run_font(fr, size=8.5, color=MUTED, italic=True)

out = r'C:\Users\muneer\dairy-simulation-prototype\client-docs\Capacity_Building_Platform_Proposal.docx'
doc.save(out)
print(out)
