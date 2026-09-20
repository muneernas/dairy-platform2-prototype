# -*- coding: utf-8 -*-
"""Generate Capacity_Building_Technical_Proposal.docx from 03-Technical-Proposal.md."""
from pathlib import Path
import re

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

ROOT = Path(__file__).resolve().parent
MD_PATH = ROOT / '03-Technical-Proposal.md'
OUT_PATH = ROOT / 'Capacity_Building_Technical_Proposal.docx'

doc = Document()

for section in doc.sections:
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

NAVY = RGBColor(0x0F, 0x3A, 0x47)
TEAL = RGBColor(0x1F, 0x6B, 0x5C)
BODY = RGBColor(0x2C, 0x2C, 0x2C)
MUTED = RGBColor(0x5A, 0x5A, 0x5A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SOFT = RGBColor(0xD5, 0xE8, 0xE2)


def set_run_font(run, name='Calibri', size=11, bold=False, color=BODY, italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def shade_cell(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'))


def set_cell_borders(cell, color='C5D5CF'):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(
        f'''<w:tcBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:left w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
        </w:tcBorders>'''
    ))


def write_cell(cell, text, bold=False, size=9, color=BODY, fill=None):
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    if fill:
        shade_cell(cell, fill)
    set_cell_borders(cell)


def clean_inline(text):
    """Remove markdown emphasis markers; keep readable punctuation."""
    text = re.sub(r'`([^`]+)`', r'\1', text)
    text = re.sub(r'\*\*([^*]+)\*\*', r'\1', text)
    text = re.sub(r'(?<!\*)\*([^*]+)\*(?!\*)', r'\1', text)
    text = text.replace('—', '—').replace('–', '–')
    text = text.replace('“', '"').replace('”', '"').replace('‘', "'").replace('’', "'")
    return text


def add_para(text, size=10.5, bold=False, color=BODY, space_before=0, space_after=5, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.12
    # **bold** segments; strip single *italics*
    text = re.sub(r'`([^`]+)`', r'\1', text)
    text = re.sub(r'(?<!\*)\*([^*]+)\*(?!\*)', r'\1', text)
    parts = re.split(r'(\*\*[^*]+\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = p.add_run(part[2:-2])
            set_run_font(run, size=size, bold=True, color=color, italic=italic)
        else:
            run = p.add_run(part)
            set_run_font(run, size=size, bold=bold, color=color, italic=italic)
    return p


def add_section_title(number, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(f'{number}  {title}')
    set_run_font(run, size=12.5, bold=True, color=NAVY)
    pPr = p._p.get_or_add_pPr()
    pPr.append(parse_xml(
        f'''<w:pBdr {nsdecls("w")}>
            <w:bottom w:val="single" w:sz="12" w:space="3" w:color="1F6B5C"/>
        </w:pBdr>'''
    ))


def add_subhead(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    set_run_font(run, size=10.5, bold=True, color=TEAL)


def add_table(headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    for i, h in enumerate(headers):
        write_cell(t.cell(0, i), h, bold=True, size=8.5, color=WHITE, fill='0F3A47')
    for r, row in enumerate(rows, 1):
        fill = 'F4F7F6' if r % 2 else 'FFFFFF'
        for c, val in enumerate(row):
            write_cell(
                t.cell(r, c),
                val,
                bold=(c == 0),
                size=8.5,
                color=NAVY if c == 0 else BODY,
                fill=fill,
            )
    return t


def strip_md(text):
    text = text.replace('\\|', '|')
    text = re.sub(r'`([^`]+)`', r'\1', text)
    text = re.sub(r'\*\*([^*]+)\*\*', r'\1', text)
    text = re.sub(r'\*([^*]+)\*', r'\1', text)
    return text.strip()


def parse_table(lines):
    """Parse markdown table lines into headers + rows."""
    rows = []
    for line in lines:
        if not line.strip().startswith('|'):
            continue
        cells = [strip_md(c) for c in line.strip().strip('|').split('|')]
        if all(re.match(r'^[-: ]+$', c) for c in cells):
            continue
        rows.append(cells)
    if not rows:
        return None, None
    return rows[0], rows[1:]


def add_banner(subject):
    banner = doc.add_table(rows=1, cols=1)
    bc = banner.cell(0, 0)
    bc.text = ''
    p1 = bc.paragraphs[0]
    r1 = p1.add_run('SMARTGREENECOS  ·  PLATFORM 2')
    set_run_font(r1, size=9, bold=True, color=WHITE)
    p2 = bc.add_paragraph()
    r2 = p2.add_run('Technical Proposal')
    set_run_font(r2, size=20, bold=True, color=WHITE)
    p3 = bc.add_paragraph()
    r3 = p3.add_run(subject[:90] if subject else 'How the platform is built')
    set_run_font(r3, size=9, color=SOFT)
    shade_cell(bc, '0F3A47')
    set_cell_borders(bc, '0F3A47')
    for p in bc.paragraphs:
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
    bc.paragraphs[0].paragraph_format.space_before = Pt(10)
    bc.paragraphs[-1].paragraph_format.space_after = Pt(10)


def add_meta(from_, date, subject):
    meta = doc.add_table(rows=3, cols=2)
    for i, (k, v) in enumerate([
        ('From', from_),
        ('Date', date),
        ('Subject', subject),
    ]):
        write_cell(meta.cell(i, 0), k, bold=True, size=9.5, color=WHITE, fill='1F6B5C')
        write_cell(meta.cell(i, 1), v, size=9.5, color=BODY, fill='F4F7F6')
    for row in meta.rows:
        row.cells[0].width = Cm(3.0)
        row.cells[1].width = Cm(14.2)


def convert_md(md_text):
    lines = md_text.splitlines()
    i = 0
    from_ = date = subject = ''
    # Skip H1, pull meta
    while i < len(lines):
        line = lines[i].strip()
        if line.startswith('# '):
            i += 1
            continue
        if line.startswith('**From:**'):
            from_ = strip_md(line.replace('**From:**', ''))
            i += 1
            continue
        if line.startswith('**Date:**'):
            date = strip_md(line.replace('**Date:**', ''))
            i += 1
            continue
        if line.startswith('**Subject:**'):
            subject = strip_md(line.replace('**Subject:**', ''))
            i += 1
            continue
        if line == '---' or line == '':
            i += 1
            if from_ and date and subject:
                break
            continue
        break

    add_banner(subject)
    add_meta(from_ or 'Investment World for Development and Technology',
             date or 'August 2026',
             subject or 'Platform 2 technical delivery')

    in_code = False
    code_buf = []
    table_buf = []

    def flush_table():
        nonlocal table_buf
        if not table_buf:
            return
        headers, rows = parse_table(table_buf)
        table_buf = []
        if headers and rows is not None:
            add_table(headers, rows)

    def flush_code():
        nonlocal code_buf
        if not code_buf:
            return
        add_para(' '.join(code_buf), size=9.5, color=MUTED, italic=True, space_after=6)
        code_buf = []

    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        stripped = line.strip()

        if stripped.startswith('```'):
            if in_code:
                flush_code()
                in_code = False
            else:
                flush_table()
                in_code = True
            i += 1
            continue

        if in_code:
            code_buf.append(stripped)
            i += 1
            continue

        if stripped.startswith('|'):
            table_buf.append(stripped)
            i += 1
            continue
        else:
            flush_table()

        if stripped == '---' or stripped == '':
            i += 1
            continue

        if stripped.startswith('*End of'):
            i += 1
            continue

        # ## 1. Title  or ## 10. Title
        m = re.match(r'^##\s+(\d+)\.\s+(.+)$', stripped)
        if m:
            num = m.group(1).zfill(2)
            add_section_title(num, m.group(2))
            i += 1
            continue

        if stripped.startswith('### '):
            add_subhead(stripped[4:])
            i += 1
            continue

        if stripped.startswith('#### '):
            add_para(strip_md(stripped[5:]), size=10, bold=True, color=TEAL, space_before=6, space_after=3)
            i += 1
            continue

        # numbered list
        if re.match(r'^\d+\.\s+', stripped):
            add_para(stripped, size=10, space_after=3)
            i += 1
            continue

        # bullet
        if stripped.startswith('- '):
            add_para('• ' + strip_md(stripped[2:]), size=10, space_after=3)
            i += 1
            continue

        # normal paragraph (may keep ** for bold rendering)
        text = stripped
        add_para(text, size=10.5, space_after=5)
        i += 1

    flush_table()
    flush_code()

    fp = doc.add_paragraph()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(14)
    fr = fp.add_run('Platform 2  ·  Technical proposal  ·  Capacity building')
    set_run_font(fr, size=8.5, color=MUTED, italic=True)


def main():
    md = MD_PATH.read_text(encoding='utf-8')
    convert_md(md)
    saved = []
    for path in (OUT_PATH, ROOT / 'Capacity_Building_Technical_Proposal_v2.docx'):
        try:
            doc.save(path)
            saved.append(str(path))
        except PermissionError:
            print(f'locked: {path}')
    for s in saved:
        print(s)
    if not saved:
        raise SystemExit('Could not save — close the Word file and retry')


if __name__ == '__main__':
    main()
